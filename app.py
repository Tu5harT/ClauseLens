"""ClauseLens - contract risk highlighter: highlights on the original PDF and fixes risky clauses."""
from __future__ import annotations

import html
import io
from pathlib import Path

import streamlit as st

import db
from risk_engine import analyze_contract, pdf_view
from risk_engine.fixer import apply_fixes
from risk_engine.rules import CATEGORY_BY_NAME

st.set_page_config(page_title="ClauseLens", page_icon="🛡️", layout="wide")

CSS = """
<style>
:root { --card:#1b2130; --card2:#222a3d; --line:#2c3650; --blue:#3b82f6; --navy:#0f2a5c; --muted:#8b95ad; }
html, body, [class*="css"], .stMarkdown, button, input, textarea { font-family: Inter, "Segoe UI", -apple-system, sans-serif; }
h1, h2, h3 { letter-spacing: -0.02em; }
.block-container { max-width: 1180px; padding-top: 1.2rem; }
#MainMenu, footer, header[data-testid="stHeader"] { visibility: hidden; }
@keyframes rise { from { opacity: 0; transform: translateY(10px); } to { opacity: 1; transform: none; } }
@keyframes sweep { from { stroke-dashoffset: var(--full); } to { stroke-dashoffset: var(--off); } }
.stTabs [data-baseweb="tab-list"] { gap: .35rem; border-bottom: 1px solid var(--line); }
.stTabs [data-baseweb="tab"] { border-radius: 10px 10px 0 0; padding: .5rem 1.1rem; color: var(--muted); }
.stTabs [aria-selected="true"] { color: #fff; background: linear-gradient(180deg, transparent, rgba(59,130,246,.18)); }
.stTabs [data-baseweb="tab-highlight"] { background: var(--blue); }
[data-testid="stFileUploader"] section { background: var(--card); border: 2px dashed #34476e; border-radius: 16px; padding: 1.6rem; transition: .2s; }
[data-testid="stFileUploader"] section:hover { border-color: var(--blue); background: var(--card2); }
[data-testid="stExpander"] { background: var(--card); border: 1px solid var(--line); border-radius: 14px; animation: rise .35s both; transition: .2s; }
[data-testid="stExpander"]:hover { border-color: #3d4d75; }
.stButton button, .stDownloadButton button { border-radius: 10px; border: 1px solid var(--line); transition: .15s; }
.stButton button:hover, .stDownloadButton button:hover { transform: translateY(-1px); border-color: var(--blue); }
.stButton button[kind="primary"], .stDownloadButton button[kind="primary"], .stFormSubmitButton button { background: linear-gradient(135deg, #1d4ed8, #3b82f6); border: 0; color: #fff; }
.hero { background: radial-gradient(900px 300px at 85% -20%, rgba(59,130,246,.35), transparent), linear-gradient(135deg, #0b1d42, #151b29 70%);
  border: 1px solid var(--line); border-radius: 20px; padding: 2rem 2.2rem; margin: .6rem 0 1.2rem; animation: rise .5s both; }
.hero h1 { font-size: 2.2rem; margin: 0 0 .3rem; }
.hero p { color: #aab4cc; margin: 0; max-width: 640px; }
.badge { display: inline-block; font-size: .72rem; letter-spacing: .08em; text-transform: uppercase; color: #8fb4ff;
  background: rgba(59,130,246,.15); border: 1px solid rgba(59,130,246,.4); padding: .2rem .6rem; border-radius: 999px; margin-bottom: .8rem; }
.card { background: var(--card); border: 1px solid var(--line); border-radius: 16px; padding: 1rem 1.2rem; animation: rise .45s both; transition: .2s; }
.card:hover { transform: translateY(-2px); border-color: #3d4d75; box-shadow: 0 8px 24px rgba(0,0,0,.25); }
.card .big { font-size: 1.9rem; font-weight: 700; line-height: 1.1; }
.card .lbl { color: var(--muted); font-size: .8rem; text-transform: uppercase; letter-spacing: .06em; }
.feat { display: flex; gap: .8rem; align-items: flex-start; margin-bottom: .7rem; }
.feat .ic { font-size: 1.4rem; background: rgba(59,130,246,.15); border-radius: 12px; padding: .45rem .6rem; }
.feat b { display: block; } .feat span { color: var(--muted); font-size: .88rem; }
.bar-row { display: flex; align-items: center; gap: .7rem; margin: .45rem 0; font-size: .86rem; }
.bar-row .name { width: 210px; color: #cfd6e6; } .bar-row .track { flex: 1; background: #10141c; border-radius: 99px; height: 9px; overflow: hidden; }
.bar-row .fill { height: 100%; border-radius: 99px; background: linear-gradient(90deg, #1d4ed8, #60a5fa); }
.bar-row .n { width: 24px; text-align: right; color: var(--muted); }
.chip { display: inline-block; font-size: .76rem; padding: .15rem .6rem; border-radius: 999px; margin: 0 .3rem .3rem 0; background: #10141c; border: 1px solid var(--line); color: #cfd6e6; }
.chip.High { border-color: #ef4444; color: #fca5a5; } .chip.Medium { border-color: #f59e0b; color: #fcd34d; } .chip.Low { border-color: #22c55e; color: #86efac; }
.quote { border-left: 3px solid var(--blue); background: #10141c; padding: .7rem 1rem; border-radius: 0 10px 10px 0; color: #d5dbea; margin-bottom: .6rem; }
.pdfpage img { border-radius: 8px; box-shadow: 0 6px 22px rgba(0,0,0,.45); }
.legend span { display: inline-block; width: 12px; height: 12px; border-radius: 3px; margin-right: 6px; vertical-align: middle; }
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

SAMPLES = Path(__file__).parent / "samples"
LEVEL_COLOR = {"High": "#ef4444", "Medium": "#f59e0b", "Low": "#22c55e"}
MARK_BG = {"High": "rgba(239,68,68,.35)", "Medium": "rgba(245,158,11,.35)", "Low": "rgba(34,197,94,.30)"}
ICONS = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
MAX_CHARS = 200_000


def extract_text(name: str, data: bytes) -> str:
    name = name.lower()
    if name.endswith(".pdf"):
        return pdf_view.extract_text(data)
    if name.endswith(".docx"):
        from docx import Document

        return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    return data.decode("utf-8", errors="ignore")


@st.cache_data(max_entries=20, ttl=3600, show_spinner=False)
def run_analysis(text: str):
    return analyze_contract(text)


@st.cache_data(max_entries=5, ttl=3600, show_spinner="Highlighting the PDF...")
def render_pdf(pdf: bytes, text: str):
    return pdf_view.annotate(pdf, run_analysis(text))


# ------------------------------------------------------------ html helpers
def gauge(score: int, level: str, caption: str = "Overall risk", size: int = 190) -> str:
    r, full = 54, 2 * 3.14159 * 54
    off = full * (1 - score / 100)
    c = LEVEL_COLOR[level]
    return (
        f'<div class="card" style="text-align:center"><svg width="{size}" height="{size}" viewBox="0 0 140 140">'
        f'<circle cx="70" cy="70" r="{r}" fill="none" stroke="#10141c" stroke-width="12"/>'
        f'<circle cx="70" cy="70" r="{r}" fill="none" stroke="{c}" stroke-width="12" stroke-linecap="round" '
        f'transform="rotate(-90 70 70)" stroke-dasharray="{full:.1f}" stroke-dashoffset="{off:.1f}" '
        f'style="--full:{full:.1f};--off:{off:.1f};animation:sweep 1.1s ease-out both"/>'
        f'<text x="70" y="68" text-anchor="middle" font-size="30" font-weight="700" fill="#fff">{score}</text>'
        f'<text x="70" y="88" text-anchor="middle" font-size="11" fill="{c}" font-weight="600">{level.upper()}</text></svg>'
        f'<div class="lbl">{caption}</div></div>'
    )


def stat_card(label: str, value, color: str = "#fff") -> str:
    return f'<div class="card"><div class="lbl">{label}</div><div class="big" style="color:{color}">{value}</div></div>'


def category_bars(counts: dict) -> str:
    if not counts:
        return ""
    top = max(counts.values())
    rows = "".join(
        f'<div class="bar-row"><div class="name">{html.escape(k)}</div><div class="track">'
        f'<div class="fill" style="width:{100 * v / top:.0f}%"></div></div><div class="n">{v}</div></div>'
        for k, v in sorted(counts.items(), key=lambda kv: -kv[1])
    )
    return f'<div class="card"><div class="lbl" style="margin-bottom:.4rem">Flagged clauses by risk type</div>{rows}</div>'


def heat_strip(report) -> str:
    """One block per clause in document order: see where the danger sits at a glance."""
    blocks = "".join(
        f'<div title="{html.escape(f"Clause {i + 1}: {f.category} ({f.level})" if f else f"Clause {i + 1}", quote=True)}" '
        f'style="flex:1;min-width:3px;height:22px;background:{LEVEL_COLOR[f.level] if f else "#2c3650"}"></div>'
        for i, (_, f) in enumerate(report.clauses)
    )
    return f'<div style="display:flex;gap:1px;border-radius:8px;overflow:hidden;margin:.3rem 0 1rem">{blocks}</div>'


def highlighted_html(report) -> str:
    parts = []
    for clause, f in report.clauses:
        safe = html.escape(clause)
        if f:
            tip = html.escape(f"{f.category} - {f.level}: {f.explanation}", quote=True)
            parts.append(
                f'<mark title="{tip}" style="background:{MARK_BG[f.level]};border-bottom:2px solid {LEVEL_COLOR[f.level]};'
                f'padding:2px 3px;border-radius:3px;color:#e6e9ef">{safe}<sup style="margin-left:4px;font-size:.65em;'
                f'font-weight:700;color:{LEVEL_COLOR[f.level]}">#{f.index + 1} {html.escape(f.category)}</sup></mark>'
            )
        else:
            parts.append(safe)
    return '<div class="card" style="line-height:2">' + " ".join(parts) + "</div>"


def docx_bytes(text: str) -> bytes:
    from docx import Document

    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ auth
def auth_forms() -> None:
    tab_in, tab_up = st.tabs(["Sign in", "Create account"])
    with tab_in, st.form("login"):
        u, p = st.text_input("Username"), st.text_input("Password", type="password")
        if st.form_submit_button("Sign in", use_container_width=True):
            user = db.verify_user(u, p)
            if user:
                st.session_state.user = user
                st.rerun()
            st.error("Wrong username or password.")
    with tab_up, st.form("register"):
        u = st.text_input("Choose a username")
        p = st.text_input("Choose a password (8+ characters)", type="password")
        if st.form_submit_button("Create account", use_container_width=True):
            err = db.create_user(u, p)
            if err:
                st.error(err)
            else:
                st.session_state.user = u.strip().lower()
                st.rerun()


FEATURES = [
    ("📑", "Highlights on your real PDF", "Risky sentences are marked right on the original pages, numbered so you know where to look."),
    ("🎯", "A risk score you can trust", "Eleven risk types, three severity levels, and a plain-English reason for every flag."),
    ("🛠", "Fix it in one click", "Swap in fairer wording, re-score instantly and download the revised contract."),
    ("🗂", "Your documents, saved", "Every analysis lives in your account so you can reopen it any time."),
]

if "user" not in st.session_state:
    left, right = st.columns([1.25, 1], gap="large")
    with left:
        st.markdown(
            '<div class="hero"><span class="badge">NLP · Contract review</span><h1>🛡️ ClauseLens</h1>'
            "<p>Read contracts the way a lawyer would: see every risky clause, understand why it matters, "
            "and fix it before you sign.</p></div>",
            unsafe_allow_html=True,
        )
        st.markdown(
            "".join(f'<div class="feat"><div class="ic">{i}</div><div><b>{t}</b><span>{d}</span></div></div>' for i, t, d in FEATURES),
            unsafe_allow_html=True,
        )
    with right:
        st.markdown('<div style="height:2.2rem"></div>', unsafe_allow_html=True)
        auth_forms()
    st.stop()

user = st.session_state.user
nav_l, nav_r = st.columns([5, 1])
nav_l.markdown(f"#### 🛡️ ClauseLens &nbsp;<span class='chip'>👤 {html.escape(user)}</span>", unsafe_allow_html=True)
if nav_r.button("Log out", use_container_width=True):
    st.session_state.clear()
    st.rerun()

rows = db.list_analyses(user)
avg = round(sum(r["overall_score"] for r in rows) / len(rows)) if rows else 0
st.markdown(
    f'<div class="hero"><span class="badge">Workspace</span><h1>Review a contract</h1>'
    f"<p>Upload a PDF, DOCX or text file. ClauseLens highlights risky clauses, explains them and helps you fix them.</p></div>",
    unsafe_allow_html=True,
)
if rows:
    s = st.columns(4)
    s[0].markdown(stat_card("Documents analysed", len(rows)), unsafe_allow_html=True)
    s[1].markdown(stat_card("Average risk", f"{avg}/100"), unsafe_allow_html=True)
    s[2].markdown(stat_card("High-risk clauses found", sum(r["high"] for r in rows), LEVEL_COLOR["High"]), unsafe_allow_html=True)
    s[3].markdown(stat_card("Clean documents", sum(r["overall_level"] == "Low" for r in rows), LEVEL_COLOR["Low"]), unsafe_allow_html=True)
    st.markdown('<div style="height:.8rem"></div>', unsafe_allow_html=True)


def analyse_and_save(name: str, text: str, pdf: bytes | None = None) -> None:
    text = text[:MAX_CHARS]
    if not text.strip():
        st.warning("Nothing to analyse (scanned PDFs need OCR first).")
        return
    report = run_analysis(text)
    if report.total_clauses == 0:
        st.warning("No readable clauses found.")
        return
    st.session_state.current = {"id": db.save_analysis(user, name, report, text, pdf), "name": name, "text": text, "pdf": pdf}


tab_up, tab_paste, tab_sample, tab_hist = st.tabs(["📤 Upload", "📋 Paste text", "🧪 Try a sample", f"🗂 My documents ({len(rows)})"])

with tab_up:
    up = st.file_uploader("PDF, DOCX or TXT", type=["pdf", "docx", "txt"], label_visibility="collapsed")
    if up and st.button("Analyse document", type="primary"):
        try:
            data = up.getvalue()
            analyse_and_save(up.name, extract_text(up.name, data), data if up.name.lower().endswith(".pdf") else None)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")

with tab_paste:
    pasted = st.text_area("Contract text", height=150, label_visibility="collapsed", placeholder="Paste contract text here")
    if st.button("Analyse pasted text", type="primary"):
        analyse_and_save("Pasted text", pasted)

with tab_sample:
    c1, c2 = st.columns(2)
    c1.markdown('<div class="card"><b>Risky SaaS contract</b><br><span class="lbl">Auto-renewal, unlimited liability, no refunds...</span></div>', unsafe_allow_html=True)
    c2.markdown('<div class="card"><b>Fair consulting agreement</b><br><span class="lbl">Balanced terms, should score low</span></div>', unsafe_allow_html=True)
    if c1.button("Analyse risky sample", use_container_width=True):
        analyse_and_save("Sample - risky SaaS contract", (SAMPLES / "sample_contract.txt").read_text(encoding="utf-8"))
    if c2.button("Analyse fair sample", use_container_width=True):
        analyse_and_save("Sample - fair consulting agreement", (SAMPLES / "fair_contract.txt").read_text(encoding="utf-8"))

with tab_hist:
    if not rows:
        st.info("Analysed documents are saved here so you can reopen them any time.")
    for r in rows:
        c = st.columns([4, 2, 2, 1, 1])
        c[0].markdown(f"**{r['filename']}**  \n<span class='lbl'>{r['created']} UTC</span>", unsafe_allow_html=True)
        c[1].markdown(f"<span class='chip {r['overall_level']}'>{r['overall_score']}/100 · {r['overall_level']}</span>", unsafe_allow_html=True)
        c[2].caption(f"{r['high']} high · {r['medium']} med · {r['low']} low")
        if c[3].button("Open", key=f"o{r['id']}"):
            row = db.get_analysis(user, r["id"])
            st.session_state.current = {"id": row["id"], "name": row["filename"], "text": row["text"], "pdf": row["pdf"]}
            st.rerun()
        if c[4].button("🗑", key=f"d{r['id']}", help="Delete"):
            db.delete_analysis(user, r["id"])
            if st.session_state.get("current", {}).get("id") == r["id"]:
                del st.session_state["current"]
            st.rerun()

# ---------------------------------------------------------------- results
current = st.session_state.get("current")
if not current:
    st.markdown("##### How it works")
    h = st.columns(3)
    for col, (n, t, d) in zip(h, [("1", "Upload", "Drop in a PDF, Word file or paste text."), ("2", "Review", "See risky clauses highlighted on the page, with reasons."), ("3", "Fix", "Apply fairer wording and download the revised contract.")]):
        col.markdown(f'<div class="card"><div class="big" style="color:#3b82f6">{n}</div><b>{t}</b><br><span class="lbl">{d}</span></div>', unsafe_allow_html=True)
    st.stop()

doc_id, pdf = current["id"], current["pdf"]
report = run_analysis(current["text"])
pngs, annotated, clause_pages = render_pdf(pdf, current["text"]) if pdf else ([], b"", {})

st.markdown("---")
st.markdown(f"### 📄 {html.escape(current['name'])}")
g, stats = st.columns([1, 2.4], gap="large")
g.markdown(gauge(report.overall_score, report.overall_level), unsafe_allow_html=True)
with stats:
    k = st.columns(4)
    k[0].markdown(stat_card("High", report.level_counts["High"], LEVEL_COLOR["High"]), unsafe_allow_html=True)
    k[1].markdown(stat_card("Medium", report.level_counts["Medium"], LEVEL_COLOR["Medium"]), unsafe_allow_html=True)
    k[2].markdown(stat_card("Low", report.level_counts["Low"], LEVEL_COLOR["Low"]), unsafe_allow_html=True)
    k[3].markdown(stat_card("Clauses read", report.total_clauses), unsafe_allow_html=True)
    st.markdown('<div style="height:.6rem"></div>', unsafe_allow_html=True)
    st.markdown(category_bars(report.category_counts) or '<div class="card">No risky clauses detected 🎉</div>', unsafe_allow_html=True)
st.markdown('<div class="lbl" style="margin-top:1rem">Risk heat-strip · one block per clause, in document order</div>', unsafe_allow_html=True)
st.markdown(heat_strip(report), unsafe_allow_html=True)

t_doc, t_fix, t_rev = st.tabs(["📄 Contract view", "🛠 Findings & fixes", "✅ Revised contract"])

with t_doc:
    if pdf:
        left, right = st.columns([3, 1])
        right.markdown(
            '<div class="card legend"><b>Legend</b><br><span style="background:#ef4444"></span>High<br>'
            '<span style="background:#f59e0b"></span>Medium<br><span style="background:#22c55e"></span>Low'
            '<p class="lbl" style="margin-top:.8rem">Margin numbers match <i>clause N</i> in Findings &amp; fixes.</p></div>',
            unsafe_allow_html=True,
        )
        right.download_button("⬇ Highlighted PDF", annotated, "highlighted.pdf", "application/pdf", use_container_width=True)
        with left.container(height=780):
            for png in pngs:
                st.image(png, use_container_width=True)
    else:
        st.markdown(highlighted_html(report), unsafe_allow_html=True)
        st.caption("Hover a highlight to see why it was flagged. Upload a PDF to see highlights on the original pages.")

selected: dict[int, str] = {}
with t_fix:
    if not report.findings:
        st.success("No risky clauses detected. Still have a lawyer review anything important.")
    else:
        keys = [f"use_{doc_id}_{f.index}" for f in report.findings]

        def set_all(value: bool) -> None:
            for key in keys:
                st.session_state[key] = value

        b1, b2, _ = st.columns([2, 1.4, 6])
        b1.button("✨ Apply all fixes", on_click=set_all, args=(True,), type="primary")
        b2.button("Clear", on_click=set_all, args=(False,))
        levels = [lv for lv in ("High", "Medium", "Low") if report.level_counts[lv]]
        for lv, tab in zip(levels, st.tabs([f"{ICONS[lv]} {lv} ({report.level_counts[lv]})" for lv in levels])):
            with tab:
                for f in (x for x in report.findings if x.level == lv):
                    page = f"<span class='chip'>page {clause_pages[f.index]}</span>" if f.index in clause_pages else ""
                    use_key, text_key = f"use_{doc_id}_{f.index}", f"fix_{doc_id}_{f.index}"
                    with st.expander(f"#{f.index + 1} · {f.category}"):
                        st.markdown(
                            f"<span class='chip {f.level}'>{f.level}</span><span class='chip'>{html.escape(f.category)}</span>{page}"
                            f"<div class='quote'>{html.escape(f.text)}</div>"
                            + "".join(f"<span class='chip'>⚠ {html.escape(r)}</span>" for r in f.explanation.split("; ")),
                            unsafe_allow_html=True,
                        )
                        st.markdown(f"**What to do:** {f.recommendation}")
                        wording = st.text_area("Replacement wording (edit freely)", CATEGORY_BY_NAME[f.category].counter, key=text_key)
                        if st.checkbox("Apply this fix", key=use_key):
                            selected[f.index] = wording

with t_rev:
    if not selected:
        st.info("Tick **Apply this fix** on findings (or use **Apply all fixes**) and the revised contract appears here.")
    else:
        revised, missed = apply_fixes(current["text"], [c for c, _ in report.clauses], selected)
        new_report = run_analysis(revised)
        a, arrow, b = st.columns([1, 0.3, 1])
        a.markdown(gauge(report.overall_score, report.overall_level, "Before", 150), unsafe_allow_html=True)
        arrow.markdown('<div style="text-align:center;font-size:2.4rem;margin-top:3rem;color:#3b82f6">➜</div>', unsafe_allow_html=True)
        b.markdown(gauge(new_report.overall_score, new_report.overall_level, f"After {len(selected)} fix(es)", 150), unsafe_allow_html=True)
        if missed:
            st.warning(f"Could not locate clause(s) {', '.join(f'#{i + 1}' for i in missed)} in the text, so they were left unchanged.")
        if pdf:
            st.caption("PDFs can't be edited in place, so the revised contract is a Word file built from the PDF's text (layout isn't preserved).")
        st.text_area("Revised text", revised, height=320)
        d1, d2 = st.columns(2)
        d1.download_button("⬇ Download revised (.docx)", docx_bytes(revised), "revised_contract.docx", type="primary", use_container_width=True)
        d2.download_button("⬇ Download revised (.txt)", revised, "revised_contract.txt", use_container_width=True)

st.caption("Automated screening aid, not legal advice. Rules and a small training set can miss or over-flag clauses; have a qualified lawyer review contracts before signing.")
